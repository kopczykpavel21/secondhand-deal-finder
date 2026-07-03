#!/usr/bin/env python3
"""
Convert an SDA/CIA monthly XLSX workbook into the normalized denominator CSV:

    brand,model,cohort_year,registrations

The workbook structure may vary across releases, so the script searches all
worksheets for a likely registrations table using fuzzy header matching.
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path
from typing import Iterable

import openpyxl


BRAND_PATTERNS = [r"znacka", r"značka", r"brand", r"make", r"vyrobce", r"výrobce"]
MODEL_PATTERNS = [r"model", r"typ", r"obchodni oznaceni", r"obchodní označení"]
YEAR_PATTERNS = [r"rok", r"year", r"obdobi", r"období"]
COUNT_PATTERNS = [r"pocet", r"počet", r"registrac", r"count", r"ks"]


def normalize(text: object) -> str:
    if text is None:
        return ""
    value = str(text).strip().lower()
    replacements = {
        "á": "a",
        "č": "c",
        "ď": "d",
        "é": "e",
        "ě": "e",
        "í": "i",
        "ň": "n",
        "ó": "o",
        "ř": "r",
        "š": "s",
        "ť": "t",
        "ú": "u",
        "ů": "u",
        "ý": "y",
        "ž": "z",
    }
    for source, target in replacements.items():
        value = value.replace(source, target)
    value = re.sub(r"\s+", " ", value)
    return value


def header_match(header: str, patterns: Iterable[str]) -> bool:
    return any(re.search(pattern, header) for pattern in patterns)


def row_is_mostly_empty(values: list[object]) -> bool:
    return sum(1 for value in values if value not in (None, "", " ")) <= 1


def infer_header_indices(header_row: list[object]) -> dict[str, int] | None:
    headers = [normalize(value) for value in header_row]
    mapping: dict[str, int] = {}

    for index, header in enumerate(headers):
        if not header:
            continue
        if "brand" not in mapping and header_match(header, BRAND_PATTERNS):
            mapping["brand"] = index
        elif "model" not in mapping and header_match(header, MODEL_PATTERNS):
            mapping["model"] = index
        elif "year" not in mapping and header_match(header, YEAR_PATTERNS):
            mapping["year"] = index
        elif "count" not in mapping and header_match(header, COUNT_PATTERNS):
            mapping["count"] = index

    if "brand" in mapping and "year" in mapping and "count" in mapping:
        return mapping
    return None


def looks_like_data_row(values: list[object], mapping: dict[str, int]) -> bool:
    try:
        brand = values[mapping["brand"]]
        year = values[mapping["year"]]
        count = values[mapping["count"]]
    except IndexError:
        return False

    if brand in (None, "", " "):
        return False

    try:
        year_int = int(str(year).strip())
    except Exception:
        return False

    if year_int < 1980 or year_int > 2035:
        return False

    try:
        int(float(str(count).replace(" ", "").replace(",", ".")))
    except Exception:
        return False

    return True


def extract_sheet_rows(ws: openpyxl.worksheet.worksheet.Worksheet) -> tuple[dict[str, int], list[dict[str, object]]] | None:
    rows = list(ws.iter_rows(values_only=True))
    for row_index, row in enumerate(rows[:60]):
        mapping = infer_header_indices(list(row))
        if mapping is None:
            continue

        extracted: list[dict[str, object]] = []
        for data_row in rows[row_index + 1:]:
            values = list(data_row)
            if row_is_mostly_empty(values):
                if extracted:
                    break
                continue
            if not looks_like_data_row(values, mapping):
                continue

            model_value = None
            if "model" in mapping and mapping["model"] < len(values):
                raw_model = values[mapping["model"]]
                if raw_model not in (None, "", " "):
                    model_value = str(raw_model).strip()

            extracted.append({
                "brand": str(values[mapping["brand"]]).strip(),
                "model": model_value,
                "cohort_year": int(str(values[mapping["year"]]).strip()),
                "registrations": int(float(str(values[mapping["count"]]).replace(" ", "").replace(",", "."))),
            })

        if extracted:
            return mapping, extracted

    return None


def parse_numeric_count(value: object) -> int | None:
    if value in (None, "", " ", "-", "-\xa0"):
        return None
    if isinstance(value, (int, float)):
        numeric = int(value)
        return numeric if numeric >= 0 else None
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text or text == "-":
        return None
    try:
        numeric = int(float(text))
    except Exception:
        return None
    return numeric if numeric >= 0 else None


def infer_sheet_year(ws: openpyxl.worksheet.worksheet.Worksheet) -> int | None:
    year_match = re.search(r"(20\d{2})", ws.title)
    if year_match:
        return int(year_match.group(1))

    rows = ws.iter_rows(values_only=True)
    try:
        first_row = next(rows)
    except StopIteration:
        return None

    for value in first_row:
        if hasattr(value, "year") and isinstance(value.year, int) and 2000 <= value.year <= 2100:
            return int(value.year)
        if isinstance(value, str):
            match = re.search(r"(20\d{2})", value)
            if match:
                return int(match.group(1))

    return None


def extract_brand_year_sheet_rows(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[dict[str, object]]:
    cohort_year = infer_sheet_year(ws)
    if cohort_year is None:
        return []

    rows = list(ws.iter_rows(values_only=True))
    extracted: list[dict[str, object]] = []

    for row in rows[4:]:
        if not row:
            continue

        brand = row[0]
        if brand in (None, "", " "):
            continue

        brand_text = str(brand).strip()
        if not brand_text or brand_text.lower() in {"členové sda", "clenove sda"}:
            continue

        registrations = None
        for candidate in row[1:4]:
            registrations = parse_numeric_count(candidate)
            if registrations is not None:
                break

        if registrations is None:
            continue

        extracted.append({
            "brand": brand_text,
            "model": None,
            "cohort_year": cohort_year,
            "registrations": registrations,
        })

    return extracted


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", help="Path to an SDA monthly XLSX file")
    parser.add_argument(
        "--output",
        default="data/denominators/cz_new_registrations.csv",
        help="Output CSV path",
    )
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    output_path = Path(args.output)

    if not xlsx_path.exists():
        raise SystemExit(f"Input file not found: {xlsx_path}")

    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

    chosen_sheet = None
    extracted_rows: list[dict[str, object]] = []
    for ws in workbook.worksheets:
        result = extract_sheet_rows(ws)
        if result is None:
            continue
        _, extracted_rows = result
        chosen_sheet = ws.title
        break

    if not extracted_rows:
        for ws in workbook.worksheets:
            brand_year_rows = extract_brand_year_sheet_rows(ws)
            if not brand_year_rows:
                continue
            extracted_rows.extend(brand_year_rows)
        chosen_sheet = "brand-year fallback across workbook"

    if not extracted_rows:
        available = ", ".join(ws.title for ws in workbook.worksheets)
        raise SystemExit(
            "Could not find a registrations table automatically. "
            f"Available sheets: {available}"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["brand", "model", "cohort_year", "registrations"],
        )
        writer.writeheader()
        writer.writerows(extracted_rows)

    print("SDA registrations conversion complete")
    print(f"input workbook : {xlsx_path}")
    print(f"sheet used     : {chosen_sheet}")
    print(f"rows written   : {len(extracted_rows)}")
    print(f"output csv     : {output_path}")


if __name__ == "__main__":
    main()
